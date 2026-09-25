#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Genera el archivo .docx de entrenamiento ENARM (sesión tarde, Pediatría)."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

OUTPUT = os.path.expanduser(
    "~/Documents/ENARM_Cerebro_Demiank/entrenamiento/"
    "entrenamiento_2026-07-15_tarde.docx"
)

doc = Document()

# Estilo global: Arial 12
style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(12)

# ============================================================
# TÍTULO GENERAL
# ============================================================
titulo = doc.add_heading('Entrenamiento ENARM — Pediatría (Sesión Tarde)', level=0)
titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Fecha: 15 de julio de 2026 — Sesión de tarde (3 temas)')
doc.add_paragraph('Especialidad: Pediatría')
doc.add_paragraph('Temas: Bronquiolitis · Intususcepción Intestinal · Estenosis Pilórica Hipertrófica')
doc.add_paragraph('')

# ============================================================
# Función auxiliar para illness script en tabla
# ============================================================
def illness_script(doc, componentes):
    tabla = doc.add_table(rows=len(componentes), cols=2)
    tabla.style = 'Light Grid Accent 1'
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (comp, contenido) in enumerate(componentes):
        celda0 = tabla.cell(i, 0)
        celda1 = tabla.cell(i, 1)
        p0 = celda0.paragraphs[0]
        run0 = p0.add_run(comp)
        run0.bold = True
        run0.font.size = Pt(11)
        celda1.paragraphs[0].add_run(contenido).font.size = Pt(11)
    doc.add_paragraph('')

# ============================================================
# Función auxiliar para preguntas
# ============================================================
def agregar_pregunta(doc, numero, caso, pregunta, opciones, respuesta, explicacion, sesgo):
    p = doc.add_paragraph()
    run = p.add_run(f'Pregunta {numero}')
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph(caso)
    doc.add_paragraph(f'Pregunta: {pregunta}')

    for letra, texto in opciones:
        doc.add_paragraph(f'{letra}) {texto}')

    p_resp = doc.add_paragraph()
    run_resp = p_resp.add_run(f'Respuesta correcta: {respuesta}')
    run_resp.bold = True
    run_resp.font.color.rgb = RGBColor(0x1A, 0x6B, 0x1A)

    doc.add_paragraph(f'Explicación: {explicacion}')
    doc.add_paragraph(f'Sesgo que activa: {sesgo}')
    doc.add_paragraph('')

# ============================================================
# TEMA 1: BRONQUIOLITIS
# ============================================================
doc.add_heading('Bronquiolitis', level=1)

illness_script(doc, [
    ('Epidemiología', 'Niños menores de 2 años, predominio entre 3 y 6 meses. Es la causa más frecuente de hospitalización en lactantes; el 90% de los casos hospitalizados son menores de 12 meses. Factores de riesgo: prematurez, peso neonatal <2,500 g, cardiopatía congénita, enfermedad pulmonar crónica del prematuro, nivel socioeconómico bajo en invierno.'),
    ('Curso temporal', 'Periodo de incubación del VSR de 4-6 días. Evolución de 3-7 días desde rinorrea y tos hasta respiración ruidosa con sibilancias. Fiebre leve. Formas graves pueden incluir cianosis.'),
    ('Fisiopatología', 'Inflamación, edema y necrosis de células epiteliales de bronquiolos con aumento de producción de moco y broncoespasmo. VSR causa hasta el 80% de los casos en epidemias.'),
    ('Presentación clínica', 'Tos, rinorrea, coriza, sibilancias audibles, prolongación de la fase espiratoria, retracciones intercostales y supraclaviculares, atrapamiento aéreo con hiperexpansión pulmonar. Fiebre leve. Sibilancias o estertores crepitantes finos e hiperinsuflación.'),
    ('Diagnóstico', 'Clínico (historia clínica y examen físico). No se recomiendan estudios radiológicos ni sanguíneos salvo sospecha de otro proceso. Oximetría de pulso para monitorización. Radiografía muestra hiperexpansión, aplanamiento diafragmático y horizontalización de arcos costales.'),
    ('Tratamiento', 'Soporte: monitorización de saturación, control térmico, hidratación, aspiración de secreciones, oxígeno con cánulas nasales. Solución salina hipertónica al 3% nebulizada. Salbutamol solo si historia familiar de atopia/asma/alergia. NO se recomiendan corticoides, antibióticos, antihistamínicos, ribavirina ni vasoconstrictores en bronquiolitis leve.'),
    ('Diagnósticos diferenciales', '1) Asma bronquial (mayor edad, antecedentes atópicos). 2) Neumonía adquirida en la comunidad (foco radiológico, fiebre alta). 3) Cuerpo extraño en vía aérea (historia súbita, asimetría radiológica).'),
])

doc.add_paragraph('ANTES de ver la respuesta, explica en voz alta por qué elegiste esa opción. Esto mejora el razonamiento (self-explanation).')
doc.add_paragraph('')

# 10 preguntas de Bronquiolitis
agregar_pregunta(doc, 1,
    'Lactante varón de 5 meses de edad, producto de parto prematuro a las 34 semanas de gestación, con peso al nacer de 2,100 g. Acude por cuadro de 4 días de evolución de rinorrea, tos y dificultad respiratoria progresiva. A la exploración: frecuencia respiratoria 60/min, saturación de O₂ 90%, sibilancias bilaterales, retracciones intercostales marcadas y cianosis perioral.',
    '¿Cuál es el agente etiológico más probable?',
    [('A', 'Virus parainfluenza tipo 1'),
     ('B', 'Virus sincitial respiratorio (VSR)'),
     ('C', 'Adenovirus'),
     ('D', 'Mycoplasma pneumoniae')],
    'B',
    'El VSR es el agente causal más común de bronquiolitis, produciendo hasta el 80% de los casos en epidemias. El paciente presenta el cuadro clásico en un lactante menor de 6 meses con factores de riesgo (prematurez, bajo peso). Los virus parainfluenza e influenza son causas menos frecuentes. Mycoplasma pneumoniae es una causa infrecuente.',
    'Disponibilidad (recordar el patógeno más conocido de vías respiratorias altas, pero no es el más frecuente en bronquiolitis)')

agregar_pregunta(doc, 2,
    'Lactante de 4 meses con bronquiolitis moderada. Frecuencia respiratoria 55/min, sibilancias en toda la espiración, tiraje intercostal, saturación 93%. Según la escala de valoración por puntos de la GPC mexicana, ¿cuál es el nivel de gravedad y la conducta adecuada?',
    '¿Cuál es la calificación y la conducta recomendada?',
    [('A', '1-3 puntos: tratamiento médico ambulatorio'),
     ('B', '4-7 puntos: referencia al segundo nivel de atención'),
     ('C', '8-14 puntos: bronquiolitis grave, hospitalización en UCIA'),
     ('D', '0 puntos: alta inmediata')],
    'B',
    'La escala de valoración por puntos clasifica: 1-3 puntos = bronquiolitis leve (ambulatorio); 4-7 puntos = bronquiolitis moderada (referencia al segundo nivel); 8-14 puntos = bronquiolitis grave. El paciente con FR 55, sibilancias en toda la espiración, tiraje intercostal y saturación 93% corresponde a un puntaje en rango moderado, por lo que requiere referencia al segundo nivel.',
    'Cierre prematuro (aceptar la primera opción plausible sin calcular el puntaje exacto)')

agregar_pregunta(doc, 3,
    'Lactante de 6 meses con bronquiolitis leve. El médico de primer nivel considera administrar salbutamol nebulizado.',
    'Según la GPC, ¿en qué pacientes con bronquiolitis está indicado el salbutamol?',
    [('A', 'En todos los pacientes con bronquiolitis, independientemente de antecedentes'),
     ('B', 'Únicamente en pacientes con historia familiar de atopia, asma o alergia'),
     ('C', 'En pacientes con saturación <92% sin importar antecedentes'),
     ('D', 'En pacientes con fiebre >38°C')],
    'B',
    'La GPC recomienda el salbutamol únicamente en pacientes con historia familiar de atopia, asma o alergia. No se recomienda el uso de broncodilatadores fuera de este contexto. El uso rutinario de salbutamol en bronquiolitis sin antecedentes atópicos no está indicado.',
    'Sobreconfianza (creer que el salbutamol es el tratamiento estándar para toda sibilancia)')

agregar_pregunta(doc, 4,
    '¿Qué porcentaje de los casos de bronquiolitis que requieren hospitalización corresponden a menores de 12 meses?',
    'Seleccione el dato correcto según la epidemiología de la bronquiolitis:',
    [('A', '50%'),
     ('B', '70%'),
     ('C', '90%'),
     ('D', '100%')],
    'C',
    'El 90% de los casos de bronquiolitis que requieren hospitalización son menores de 12 meses. Este dato epidemiológico es clave para el ENARM, ya que orienta la valoración del riesgo según edad.',
    'Anclaje (fijarse en el primer número redondo que parece plausible)')

agregar_pregunta(doc, 5,
    'Lactante de 3 meses con bronquiolitis que no mejora con manejo ambulatorio. El médico de primer nivel debe decidir si referencia al segundo nivel.',
    '¿Cuál de los siguientes NO es un criterio de referencia al segundo nivel de atención según la GPC?',
    [('A', 'Antecedente de prematurez'),
     ('B', 'Edad <3 meses'),
     ('C', 'Bronquiolitis moderada o grave'),
     ('D', 'Familia con recursos económicos altos')],
    'D',
    'Los criterios de referencia incluyen: antecedente de prematurez, edad <3 meses, bronquiolitis moderada o grave, comorbilidad o factores de riesgo, nivel socioeconómico familiar BAJO, padres adolescentes, hermano fallecido por enfermedad pulmonar antes de los 5 años, evolución clínica no favorable e intolerancia a la vía oral. Una familia con recursos económicos altos NO es criterio de referencia.',
    'Prevalencia base (pensar que los recursos económicos altos protegerían, pero el criterio real es lo contrario)')

agregar_pregunta(doc, 6,
    'Niño de 8 meses con bronquiolitis. El médico solicita radiografía de tórax. ¿Cuál es el hallazgo radiológico característico?',
    'Identifique el patrón radiológico típico de la bronquiolitis:',
    [('A', 'Consolidación lobar con broncograma aéreo'),
     ('B', 'Aumento de la radiolucidez, hiperexpansión pulmonar, aplanamiento diafragmático y horizontalización de arcos costales'),
     ('C', 'Derrame pleural bilateral'),
     ('D', 'Neumotórax a tensión')],
    'B',
    'La radiografía torácica en bronquiolitis muestra aumento de la radiolucidez de los campos pulmonares, hiperexpansión del parénquima pulmonar, aplanamiento o depresión de las sombras diafragmáticas y horizontalización de los arcos costales, todo secundario al atrapamiento aéreo. La consolidación lobar sugiere neumonía, no bronquiolitis.',
    'Disponibilidad (asociar radiografía de tórax con consolidación, que es lo más enseñado)')

agregar_pregunta(doc, 7,
    'Lactante de 2 meses con bronquiolitis leve. La madre pregunta si puede darle paracetamol. ¿Cuál es la indicación de paracetamol según la GPC?',
    '¿Cuándo se recomienda administrar paracetamol en bronquiolitis?',
    [('A', 'Cuando la temperatura corporal es >38°C'),
     ('B', 'Cuando la temperatura corporal es >37.5°C'),
     ('C', 'Solo si hay sibilancias graves'),
     ('D', 'Nunca se recomienda en bronquiolitis')],
    'A',
    'El paracetamol se recomienda cuando la temperatura corporal es >38°C. Es parte de las medidas de soporte (control térmico). No se indica para temperaturas menores ni como rutina.',
    'Anclaje (fijarse en 37.5°C como umbral común de fiebre, pero la GPC específica 38°C)')

agregar_pregunta(doc, 8,
    'Recién nacido de 34 semanas de gestación con displasia broncopulmonar grave. Se planea profilaxis para VSR.',
    '¿Cuál es el fármaco indicado para la prevención de cuadros graves por VSR y cuál es su esquema?',
    [('A', 'Vacuna contra VSR, dosis única intramuscular'),
     ('B', 'Palivizumab, inyecciones mensuales iniciando antes de la temporada de mayor incidencia'),
     ('C', 'Ribavirina oral diaria durante el invierno'),
     ('D', 'Oseltamivir, una dosis diaria por 10 días')],
    'B',
    'Palivizumab es un anticuerpo monoclonal específico para el VSR que se administra en inyecciones mensuales, iniciando justo antes del inicio de la temporada de mayor incidencia. Está indicado en menores de 2 años con enfermedad pulmonar crónica (displasia broncopulmonar grave), infantes de muy bajo peso y pacientes con cardiopatías congénitas cianóticas y acianóticas hemodinámicamente significativas.',
    'Cierre prematuro (pensar en vacuna como prevención estándar, pero no existe vacuna contra VSR para este grupo)')

agregar_pregunta(doc, 9,
    'Lactante de 5 meses con bronquiolitis moderada. El médico considera administrar solución salina hipertónica nebulizada.',
    '¿Cuál es la concentración de solución salina hipertónica recomendada y su mecanismo de acción?',
    [('A', 'Solución salina al 0.9% — humidifica la vía aérea'),
     ('B', 'Solución salina hipertónica al 3% — moviliza agua al intersticio de las vías aéreas, disminuye edema y viscosidad del moco'),
     ('C', 'Solución salina hipertónica al 7% — reduce directamente el broncoespasmo'),
     ('D', 'Solución salina al 5% — efecto antibacteriano directo')],
    'B',
    'La solución salina hipertónica al 3% nebulizada mejora la hidratación de las vías aéreas por movilización del agua al intersticio de las vías aéreas, con disminución del edema intersticial y de la viscosidad de la mucosa. La concentración es 3%, no 0.9% (salina normal) ni 7%.',
    'Sobreconfianza (creer que la salina normal 0.9% es la indicada por ser lo más usado)')

agregar_pregunta(doc, 10,
    'Lactante de 7 meses con cuadro de tos, rinorrea y sibilancias. El médico solicita biometría hemática, PCR y radiografía de tórax de rutina.',
    'Según la GPC, ¿cuál es la recomendación sobre estudios de laboratorio y gabinete en bronquiolitis?',
    [('A', 'Solicitar biometría hemática y PCR en todos los casos'),
     ('B', 'El diagnóstico es clínico; no se recomiendan estudios radiológicos ni sanguíneos salvo sospecha de otro proceso infeccioso'),
     ('C', 'La radiografía de tórax es obligatoria en todo caso de bronquiolitis'),
     ('D', 'Se requiere PCR para diferenciar VSR de adenovirus')],
    'B',
    'El diagnóstico de bronquiolitis es clínico, basado en la historia clínica y el examen físico. No se recomienda realizar estudios radiológicos ni sanguíneos a menos que se sospeche otro proceso infeccioso. La oximetría de pulso sí es necesaria para monitorización de la saturación de oxígeno.',
    'Cierre prematuro (aceptar que todo paciente con dificultad respiratoria necesita laboratorio y gabinete)')

# ============================================================
# TEMA 2: INTUSUSCEPCIÓN INTESTINAL
# ============================================================
doc.add_page_break()
doc.add_heading('Intususcepción Intestinal', level=1)

illness_script(doc, [
    ('Epidemiología', 'La mayoría de los casos ocurren en lactantes de 4-9 meses, principalmente en hombres (relación 4:1). Casi todos los episodios son idiopáticos en menores de 2 años. Asociación temporal con infección vírica (hiperplasia linfoide). En niños mayores aumenta la frecuencia de un punto patológico (ej. divertículo de Meckel).'),
    ('Curso temporal', 'Aparición súbita de dolor abdominal espasmódico cada 20-30 minutos, alternado con letargo (aspecto "sedado"). Progresión a vómito biliar, fatiga intestinal con disminución del dolor y heces "en jalea de grosella".'),
    ('Fisiopatología', 'Introducción telescópica de un segmento proximal (intususcepto) en el distal (intususcipiente). El avance del intususcepto comprime y estira el mesenterio, obstruye el drenaje venoso, produce edema, congestión y sangrado. Genera un tercer espacio de fluido.'),
    ('Presentación clínica', 'Dolor abdominal espasmódico súbito, elevación de rodillas, llanto con palidez, patrón cólico cada 20-30 min, letargo intercrítico, rechazo al alimento, vómito biliar, heces "en jalea de grosella", masa palpable en forma de salchicha en cuadrante superior derecho o epigastrio.'),
    ('Diagnóstico', 'Radiografía de abdomen (estudio inicial). Ultrasonido abdominal (estudio complementario). Colon por enema hidrostático o neumático (estudio de elección: diagnóstico y tratamiento).'),
    ('Tratamiento', 'Reanimación hídrica y sonda nasogástrica antes de la intervención. Reducción no quirúrgica por colon por enema si evolución <24 h, sin perforación/peritonitis/sepsis, sin oclusión alta, sin rectorragia importante, en menores de 2 años. Intervención quirúrgica urgente si no se logra reducción o hay necrosis/perforación. Riesgo de recurrencia 5-10%.'),
    ('Diagnósticos diferenciales', '1) Apendicitis aguda (dolor en fosa iliaca derecha, fiebre progresiva). 2) Vólvulo intestinal (distensión abrupta, vómito bilioso). 3) Divertículo de Meckel sangrante (sangrado indoloro, dolor en fosa iliaca derecha).'),
])

doc.add_paragraph('ANTES de ver la respuesta, explica en voz alta por qué elegiste esa opción. Esto mejora el razonamiento (self-explanation).')
doc.add_paragraph('')

agregar_pregunta(doc, 1,
    'Lactante varón de 6 meses, previamente sano, que acude por episodios de llanto intenso con palidez y elevación de rodillas que ocurren cada 25 minutos, alternados con periodos de letargo. La madre refiere un episodio de vómito no bilioso. A la exploración se palpa una masa en el cuadrante superior derecho.',
    '¿Cuál es el diagnóstico más probable?',
    [('A', 'Apendicitis aguda'),
     ('B', 'Intususcepción intestinal'),
     ('C', 'Estenosis pilórica hipertrófica'),
     ('D', 'Divertículo de Meckel')],
    'B',
    'El cuadro clásico de intususcepción intestinal en un lactante de 4-9 meses incluye dolor abdominal espasmódico cada 20-30 minutos, elevación de rodillas, llanto con palidez, letargo intercrítico y masa palpable en forma de salchicha en cuadrante superior derecho. La apendicitis es más frecuente en mayores de 2 años. La estenosis pilórica presenta vómito en proyectil no bilioso sin dolor cólico. El divertículo de Meckel generalmente cursa con sangrado indoloro.',
    'Disponibilidad (pensar en apendicitis como causa más común de abdomen agudo pediátrico)')

agregar_pregunta(doc, 2,
    'Lactante de 7 meses con sospecha de intususcepción intestinal. El cirujano evalúa los criterios para reducción no quirúrgica por colon por enema.',
    '¿Cuál de los siguientes es un criterio de indicación para descompresión por colon por enema?',
    [('A', 'Evolución mayor a 48 horas'),
     ('B', 'Presencia de peritonitis'),
     ('C', 'Evolución menor a 24 horas'),
     ('D', 'Rectorragia importante')],
    'C',
    'Los criterios de indicación para descompresión por colon por enema (hidrostático o neumático) son: 1) evolución menor a 24 horas, 2) ausencia de datos de perforación, peritonitis o sepsis, 3) ausencia de oclusión intestinal alta, 4) ausencia de rectorragia importante, 5) niños menores de 2 años. La evolución mayor a 48 horas, la peritonitis y la rectorragia importante son contraindicaciones.',
    'Cierre prematuro (aceptar la primera opción que parece clínica sin revisar los criterios exactos)')

agregar_pregunta(doc, 3,
    'Lactante de 8 meses con intususcepción intestinal confirmada. Se realiza reducción hidrostática exitosa.',
    '¿Cuál es el riesgo de recurrencia después de una reducción exitosa y por cuánto tiempo debe vigilarse al paciente?',
    [('A', '1-2% durante 12 horas'),
     ('B', '5-10% durante 24 horas'),
     ('C', '20-30% durante 48 horas'),
     ('D', 'No hay riesgo de recurrencia')],
    'B',
    'Después de una reducción completa exitosa, el paciente debe ser vigilado durante 24 horas, ya que existe un riesgo de recurrencia del 5-10%. Si la reducción no es completa o hay datos de necrosis o perforación, se requiere intervención quirúrgica urgente.',
    'Anclaje (fijarse en el dato de 24 horas pero confundir el porcentaje de recurrencia)')

agregar_pregunta(doc, 4,
    'Lactante de 5 meses con intususcepción intestinal. Se realiza una radiografía de abdomen como estudio inicial.',
    '¿Cuál es el orden correcto de los estudios diagnósticos en intususcepción intestinal?',
    [('A', 'Colon por enema → radiografía → ultrasonido'),
     ('B', 'Radiografía de abdomen → ultrasonido abdominal → colon por enema'),
     ('C', 'Ultrasonido → tomografía → colon por enema'),
     ('D', 'Tomografía abdominal → radiografía → laparoscopia')],
    'B',
    'El orden de los estudios es: 1) Radiografía de abdomen (estudio inicial), 2) Ultrasonido abdominal (estudio complementario, forma simple de llegar al diagnóstico), 3) Colon por enema hidrostático o neumático (estudio de elección, diagnóstico y tratamiento). El colon por enema es el estudio más directo y potencialmente útil.',
    'Sobreconfianza (creer que el ultrasonido es el primer estudio por ser no invasivo)')

agregar_pregunta(doc, 5,
    'Lactante de 9 meses, sexo masculino, con dolor abdominal cólico intermitente y masa palpable en epigastrio. Las heces tienen aspecto de "jalea de grosella".',
    '¿Cuál es el tipo anatómico más frecuente de intususcepción en menores de 2 años?',
    [('A', 'Ileoileal'),
     ('B', 'Yeyunoyeyunal'),
     ('C', 'Ileocecocólica'),
     ('D', 'Colocolónica')],
    'C',
    'El tipo más frecuente de intususcepción en menores de 2 años es la ileocecocólica. En niños mayores, donde es más frecuente encontrar un punto patológico, se observan también intususcepciones ileoileales, yeyunoileales y yeyunoyeyunales.',
    'Prevalencia base (salirse a las formas menos frecuentes por parecer más específicas)')

agregar_pregunta(doc, 6,
    'Lactante de 6 meses con intususcepción intestinal. ¿Cuál es la relación hombre:mujer reportada?',
    'Identifique el dato epidemiológico correcto:',
    [('A', '1:1'),
     ('B', '2:1'),
     ('C', '4:1'),
     ('D', '1:4')],
    'C',
    'La intususcepción intestinal ocurre principalmente en hombres con una relación 4:1 respecto a las mujeres. Este dato epidemiológico, junto con el rango de edad de 4-9 meses, es característico de la presentación típica idiopática.',
    'Anclaje (fijarse en 2:1 como proporción común en patología pediátrica)')

agregar_pregunta(doc, 7,
    'Lactante de 5 meses con intususcepción intestinal de 12 horas de evolución. Se decide manejo no quirúrgico.',
    '¿Cuál es el tratamiento inicial antes de la reducción radiológica?',
    [('A', 'Antibióticos de amplio espectro intravenosos'),
     ('B', 'Reanimación hídrica y colocación de sonda nasogástrica'),
     ('C', 'Analgesia con opioides sin hidratación'),
     ('D', 'Cirugía exploradora inmediata')],
    'B',
    'El tratamiento inicial se realiza con reanimación hídrica y colocación de sonda nasogástrica, para lograr la reposición hídrica completa antes de la intervención radiológica. La evaluación quirúrgica se solicita con frecuencia para que el cirujano esté presente durante la reducción no quirúrgica. Los antibióticos no son el manejo inicial estándar.',
    'Cierre prematuro (pensar en antibióticos como primera medida por asociar con infección)')

agregar_pregunta(doc, 8,
    'Niño de 4 años con intususcepción intestinal recurrente. El médico sospecha un punto patológico subyacente.',
    '¿Cuál es la causa más probable de un punto patológico en niños mayores con intususcepción?',
    [('A', 'Hiperplasia linfoide idiopática'),
     ('B', 'Divertículo de Meckel'),
     ('C', 'Tumor de Wilms'),
     ('D', 'Pólipos colónicos múltiples')],
    'B',
    'En niños mayores aumenta la frecuencia en la identificación de un punto patológico, siendo el divertículo de Meckel el ejemplo mencionado en la nota. En menores de 2 años, casi todos los episodios son idiopáticos y se asocian con hiperplasia linfoide por infección vírica.',
    'Disponibilidad (pensar en hiperplasia linfoide por ser la causa más conocida en lactantes)')

agregar_pregunta(doc, 9,
    'Lactante de 7 meses con intususcepción intestinal. La madre refiere que el niño tiene episodios de letargo intercrítico y "parece sedado".',
    '¿Cuál es el mecanismo fisiopatológico que explica el sangrado en heces "en jalea de grosella"?',
    [('A', 'Úlcera péptica por estrés gástrico'),
     ('B', 'Compresión y estiramiento del mesenterio con obstrucción del drenaje venoso, edema, congestión y sangrado'),
     ('C', 'Infección bacteriana invasiva del colon'),
     ('D', 'Isquemia arterial aguda por embolia')],
    'B',
    'El avance del intususcepto provoca compresión y estiramiento del mesenterio, con obstrucción del drenaje venoso, edema y congestión con sangrado. Esto lleva a la generación de un tercer espacio de fluido y a las heces "en jalea de grosella". No es un proceso infeccioso ni una úlcera péptica.',
    'Sobreconfianza (atribuir el sangrado a infección por ser un mecanismo común en pediatría)')

agregar_pregunta(doc, 10,
    'Lactante de 6 meses con intususcepción intestinal. Tras la reducción por enema, el cirujano solicita vigilancia.',
    '¿Cuál de las siguientes es una indicación de intervención quirúrgica urgente tras intentar la reducción no quirúrgica?',
    [('A', 'Reducción completa con buena tolerancia oral'),
     ('B', 'Reducción incompleta o datos de necrosis/perforación'),
     ('C', 'Lloro intenso post-reducción'),
     ('D', 'Fiebre transitoria post-reducción')],
    'B',
    'Si se logra la reducción completa, el paciente se vigila 24 horas. Si la reducción no es completa, o se presentan datos de necrosis y/o perforación, se requiere intervención quirúrgica urgente. El lloro intenso y la fiebre transitoria no son indicaciones quirúrgicas per se.',
    'Cierre prematuro (aceptar la fiebre como signo de complicación sin evaluar el contexto)')

# ============================================================
# TEMA 3: ESTENOSIS PILÓRICA HIPERTRÓFICA
# ============================================================
doc.add_page_break()
doc.add_heading('Estenosis Pilórica Hipertrófica', level=1)

illness_script(doc, [
    ('Epidemiología', 'Ocurre en 1.5-4 por cada 1,000 nacidos vivos. Predominio en sexo masculino 4-5:1. Factores de riesgo: sexo masculino, primogenitura, antecedente familiar, hemotipos B y O, tratamiento con eritromicina en las primeras 2 semanas de vida, hijos de mujeres tratadas con macrólidos en embarazo o lactancia. Asociación con fístula traqueoesofágica, trisomía 18, síndromes de Apert, Zellweger, Smith-Lemli-Opitz y Cornelia de Lange.'),
    ('Curso temporal', 'Se presenta entre las 2-8 semanas de vida (rango 1 semana a 3 meses), más tardía en prematuros. Vómito no biliar progresivo, postprandial, profuso, en proyectil. Evolución a letargo, pérdida de peso y deshidratación.'),
    ('Fisiopatología', 'Hipertrofia e hiperplasia de la capa muscular de la porción antro-pilórica del estómago, que disminuye la luz intestinal y produce obstrucción al vaciamiento gástrico. Es la causa más frecuente de cirugía abdominal en menores de 2 meses y la causa más común de alcalosis metabólica en pediatría.'),
    ('Presentación clínica', 'Vómito no biliar, progresivo, postprandial, profuso, en proyectil. Apetito voraz que evoluciona a letargo, pérdida de peso y deshidratación. Abultamiento gástrico, peristaltismo gástrico visible. Palpación de la "oliva pilórica" (signo patognomónico). Puede acompañarse de ictericia leve.'),
    ('Diagnóstico', 'Ultrasonido (primera elección, sensibilidad 98%, especificidad 100%). Serie esófago-gastroduodenal (segunda elección, si ultrasonido no concluyente). Endoscopia (tercera elección). Si se palpa la oliva pilórica, no se requiere ningún estudio de gabinete.'),
    ('Tratamiento', 'Reposición hidroelectrolítica prequirúrgica (corregir alcalosis metabólica hipoclorémica e hipocalémica). Cirugía: piloromiotomía de Fredet-Ramstedt por abordaje laparoscópico. Complicación transquirúrgica principal: perforación de la mucosa pilórica. Analgesia posoperatoria: paracetamol.'),
    ('Diagnósticos diferenciales', '1) Reflujo gastroesofágico (vómito no proyectil, no progresivo). 2) Estenosis duodenal/atresia duodenal (vómito bilioso). 3) Sepsis neonatal (letargo, fiebre, hallazgos sistémicos).'),
])

doc.add_paragraph('ANTES de ver la respuesta, explica en voz alta por qué elegiste esa opción. Esto mejora el razonamiento (self-explanation).')
doc.add_paragraph('')

agregar_pregunta(doc, 1,
    'Lactante varón de 4 semanas de vida, producto de primer embarazo, que acude por vómito postprandial no bilioso, progresivo, en proyectil, de 5 días de evolución. El niño se muestra con apetito voraz después de cada episodio. A la exploración se palpa una masa en el epigastrio.',
    '¿Cuál es el diagnóstico más probable?',
    [('A', 'Reflujo gastroesofágico'),
     ('B', 'Estenosis pilórica hipertrófica'),
     ('C', 'Atresia duodenal'),
     ('D', 'Sepsis neonatal')],
    'B',
    'El cuadro clásico de estenosis pilórica hipertrófica incluye vómito no biliar, progresivo, postprandial, en proyectil, en un lactante de 2-8 semanas de vida, con apetito voraz y palpación de la "oliva pilórica". El reflujo gastroesofágico produce vómito no proyectil. La atresia duodenal presenta vómito bilioso. La sepsis neonatal tiene un cuadro sistémico.',
    'Anclaje (fijarse en "vómito" y pensar en reflujo gastroesofágico por ser más común)')

agregar_pregunta(doc, 2,
    'Lactante de 3 semanas con vómito en proyectil no bilioso. El médico palpa la oliva pilórica en la exploración física.',
    'Según la GPC, ¿cuál es la conducta diagnóstica adecuada cuando se palpa la oliva pilórica?',
    [('A', 'Solicitar ultrasonido de manera obligatoria para confirmar'),
     ('B', 'Solicitar serie esófago-gastroduodenal'),
     ('C', 'No es necesario realizar ningún estudio de gabinete; la palpación de la oliva es patognomónica'),
     ('D', 'Solicitar biometría hemática y electrolitos antes de cualquier estudio')],
    'C',
    'La GPC establece que la palpación de la oliva pilórica es el signo patognomónico y que, si esta se palpa, no es necesario realizar ningún estudio de gabinete. El diagnóstico es clínico en este escenario. Los estudios de imagen se reservan para casos donde no se palpa la oliva.',
    'Cierre prematuro (aceptar que siempre se necesita confirmación radiológica)')

agregar_pregunta(doc, 3,
    'Lactante de 5 semanas con estenosis pilórica hipertrófica confirmada por ultrasonido. Se planea tratamiento quirúrgico.',
    '¿Cuál es el tratamiento quirúrgico de elección?',
    [('A', 'Piloroplastia con vagotomía'),
     ('B', 'Piloromiotomía de Fredet-Ramstedt por abordaje laparoscópico'),
     ('C', 'Gastroyeyunostomía'),
     ('D', 'Resección del antro pilórico con anastomosis')],
    'B',
    'El tratamiento quirúrgico de elección es la piloromiotomía de Fredet-Ramstedt, actualmente por abordaje laparoscópico. Se realiza tras la reposición hidroelectrolítica prequirúrgica. La piloroplastia con vagotomía, la gastroyeyunostomía y la resección no son los procedimientos indicados.',
    'Sobreconfianza (pensar en piloroplastia por ser un procedimiento similar en nombre)')

agregar_pregunta(doc, 4,
    'Lactante de 4 semanas con vómito en proyectil de 10 días de evolución. Se solicitan electrolitos séricos.',
    '¿Cuál es el trastorno ácido-base característico de la estenosis pilórica hipertrófica?',
    [('A', 'Acidosis metabólica hiperclorémica'),
     ('B', 'Alcalosis metabólica hipoclorémica e hipocalémica'),
     ('C', 'Acidosis respiratoria no compensada'),
     ('D', 'Alcalosis respiratoria aguda')],
    'B',
    'La estenosis pilórica hipertrófica es la causa más común de alcalosis metabólica en pediatría. Produce alcalosis metabólica hipoclorémica e hipocalémica por pérdida de ácido clorhídrico y potasio por el vómito. Puede acompañarse de elevación del nitrógeno ureico e hiperbilirrubinemia no conjugada.',
    'Anclaje (fijarse en "acidosis metabólica" por ser el trastorno más enseñado en general)')

agregar_pregunta(doc, 5,
    'Lactante de 3 semanas con sospecha de estenosis pilórica. No se palpa la oliva pilórica. Se solicita ultrasonido.',
    '¿Cuáles son los criterios ecográficos diagnósticos de la estenosis pilórica hipertrófica?',
    [('A', 'Diámetro pilórico total >10 mm, espesor del músculo >2 mm, longitud >10 mm'),
     ('B', 'Diámetro pilórico total >15-18 mm, espesor del músculo >3-4 mm, longitud >17 mm'),
     ('C', 'Diámetro pilórico total >20 mm, espesor del músculo >5 mm, longitud >25 mm'),
     ('D', 'Diámetro pilórico total >12 mm, espesor del músculo >3 mm, longitud >15 mm')],
    'B',
    'Los criterios ecográficos de la GPC son: diámetro pilórico total >15-18 mm, espesor del músculo pilórico >3-4 mm, longitud del canal pilórico >17 mm, con imagen característica de "doble riel" por el estrechamiento de la luz intestinal. El ultrasonido tiene sensibilidad del 98% y especificidad del 100%.',
    'Cierre prematuro (aceptar el primer conjunto de cifras que parecen plausibles)')

agregar_pregunta(doc, 6,
    'Recién nacido tratado con eritromicina durante las primeras 2 semanas de vida por sospecha de conjuntivitis.',
    '¿Cuál es el factor de riesgo asociado a estenosis pilórica hipertrófica identificado en este paciente?',
    [('A', 'Tratamiento con eritromicina durante las primeras 2 semanas de vida'),
     ('B', 'Tratamiento con amoxicilina durante las primeras 2 semanas de vida'),
     ('C', 'Tratamiento con paracetamol durante las primeras 2 semanas de vida'),
     ('D', 'Tratamiento con vitamina K durante las primeras 2 semanas de vida')],
    'A',
    'El tratamiento con eritromicina durante las primeras 2 semanas de vida extrauterina es un factor de riesgo alto para el desarrollo de estenosis pilórica hipertrófica. También son factores de riesgo: hijos de mujeres tratadas con macrólidos durante el embarazo o la lactancia, sexo masculino, primogenitura, antecedente familiar y hemotipos B y O.',
    'Disponibilidad (pensar en amoxicilina por ser el antibiótico más común en pediatría)')

agregar_pregunta(doc, 7,
    'Lactante de 6 semanas con estenosis pilórica confirmada. Se decide realizar serie esófago-gastroduodenal porque el ultrasonido no fue concluyente.',
    '¿Cuál es el hallazgo característico de la serie esófago-gastroduodenal en estenosis pilórica?',
    [('A', 'Imagen de "doble o triple riel" por elongación y engrosamiento del conducto pilórico'),
     ('B', 'Imagen de "colon en cortes"'),
     ('C', 'Signo de la "burbuja doble"'),
     ('D', 'Signo del "asa centinela"')],
    'A',
    'Los hallazgos característicos de la serie esófago-gastroduodenal incluyen: gastromegalia, retraso en el vaciamiento gástrico, onda antiperistáltica que se detiene en el estómago, elongación (2-3 cm) del conducto pilórico con engrosamiento que produce una imagen de "doble o triple riel", y un signo de acúmulo de bario en el antro pilórico dilatado. La "burbuja doble" es signo de atresia duodenal, no de estenosis pilórica.',
    'Anclaje (fijarse en "burbuja doble" por ser un signo radiológico conocido en pediatría)')

agregar_pregunta(doc, 8,
    'Lactante de 4 semanas con estenosis pilórica hipertrófica. Se programa para piloromiotomía de Fredet-Ramstedt.',
    '¿Cuál es la principal complicación transquirúrgica de la piloromiotomía?',
    [('A', 'Hemorragia masiva del lecho pilórico'),
     ('B', 'Perforación de la mucosa pilórica'),
     ('C', 'Lesión del nervio vago'),
     ('D', 'Fístula gastrocutánea')],
    'B',
    'La principal complicación transquirúrgica de la piloromiotomía de Fredet-Ramstedt es la perforación de la mucosa pilórica, que se repara con cierre primario, efectuándose la piloromiotomía en otro punto de la circunferencia del píloro. Como analgesia posoperatoria se recomienda el paracetamol como fármaco de elección.',
    'Cierre prematuro (pensar en hemorragia por ser la complicación más temida en cirugía en general)')

agregar_pregunta(doc, 9,
    'Lactante de 5 semanas con vómito en proyectil. Se solicita ultrasonido, que muestra un canal pilórico de 20 mm de longitud y un espesor muscular de 4 mm.',
    '¿Cuál es la sensibilidad y especificidad del ultrasonido en el diagnóstico de estenosis pilórica hipertrófica?',
    [('A', 'Sensibilidad 80%, especificidad 90%'),
     ('B', 'Sensibilidad 98%, especificidad 100%'),
     ('C', 'Sensibilidad 90%, especificidad 95%'),
     ('D', 'Sensibilidad 100%, especificidad 98%')],
    'B',
    'El ultrasonido es el estudio de primera elección con sensibilidad del 98% y especificidad del 100%. Los criterios ecográficos diagnósticos son: diámetro pilórico total >15-18 mm, espesor del músculo pilórico >3-4 mm, longitud del canal pilórico >17 mm, con imagen de "doble riel".',
    'Sobreconfianza (creer que la especificidad no puede ser 100% por ser estadísticamente "demasiado perfecta")')

agregar_pregunta(doc, 10,
    'Lactante de 3 semanas con vómito en proyectil no bilioso. Se diagnostica estenosis pilórica hipertrófica y se programa para cirugía.',
    '¿Cuál es la causa principal de retraso de la cirugía en pacientes con estenosis pilórica hipertrófica?',
    [('A', 'Falta de consentimiento informado de los padres'),
     ('B', 'Desequilibrio hidroelectrolítico no corregido'),
     ('C', 'Fiebre postoperatoria anticipada'),
     ('D', 'Ausencia de quirófano disponible')],
    'B',
    'La reposición hidroelectrolítica prequirúrgica es esencial con la finalidad de evitar complicaciones metabólicas. El desequilibrio hidroelectrolítico (alcalosis metabólica hipoclorémica e hipocalémica) es una de las principales causas de retraso de la cirugía. La cirugía se realiza tras la corrección completa de los electrolitos.',
    'Cierre prematuro (aceptar la falta de quirófano como causa más frecuente sin considerar el aspecto metabólico)')

# ============================================================
# Guardar
# ============================================================
os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
doc.save(OUTPUT)
print(f'Archivo guardado: {OUTPUT}')